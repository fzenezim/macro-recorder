# -*- coding: utf-8 -*-
"""Testes do excel_data: extract, make_template, load_rows."""
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from macro_recorder.events import Step, Action
from macro_recorder.excel_data import (
    extract_placeholders,
    make_template,
    load_rows,
)


def _type(text: str, redacted: bool = False) -> Step:
    return Step(action=Action.TYPE.value, t=1.0, text=text, redacted=redacted)


def _click() -> Step:
    return Step(action=Action.CLICK.value, t=2.0, x=5, y=6)


# ── extract_placeholders ─────────────────────────────────────────
def test_extract_single():
    assert extract_placeholders([_type("Olá {nome}")]) == ["nome"]


def test_extract_ordered_unique():
    steps = [_type("{a} {b} {a}"), _type("x{c}x")]
    assert extract_placeholders(steps) == ["a", "b", "c"]


def test_extract_ignores_non_type():
    assert extract_placeholders([_click(), _type("fixo")]) == []


def test_extract_ignores_redacted():
    # text real ainda, mas redacted=True → não contribui
    assert extract_placeholders([_type("{senha}", redacted=True)]) == []


def test_extract_no_placeholder():
    assert extract_placeholders([_type("sem nada")]) == []


# ── make_template ───────────────────────────────────────────────
def test_template_headers_order(tmp_path: Path):
    steps = [_type("cpf {cpf}"), _type("nome {nome} {cpf}")]
    out = make_template(steps, tmp_path / "data.xlsx")
    assert out.exists()
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["cpf", "nome"]


def test_template_no_placeholder_generic(tmp_path: Path):
    out = make_template([_type("fixo")], tmp_path / "data.xlsx")
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["campo1"]


def test_template_bold_header(tmp_path: Path):
    make_template([_type("{a}")], tmp_path / "data.xlsx")
    ws = load_workbook(out if (out := tmp_path / "data.xlsx") else tmp_path / "x.xlsx").active
    assert ws[1][0].font.b


def test_template_empty_data_row(tmp_path: Path):
    make_template([_type("{a}")], tmp_path / "data.xlsx")
    assert load_rows(tmp_path / "data.xlsx") == []


# ── load_rows ───────────────────────────────────────────────────
def test_load_basic(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["nome", "valor"])
    ws.append(["a", 1])
    ws.append(["b", 2])
    f = tmp_path / "x.xlsx"
    wb.save(f)
    assert load_rows(f) == [
        {"nome": "a", "valor": "1"},
        {"nome": "b", "valor": "2"},
    ]


def test_load_skips_blank_row(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["k"])
    ws.append(["v1"])
    ws.append([None])
    ws.append(["v2"])
    f = tmp_path / "x.xlsx"
    wb.save(f)
    assert load_rows(f) == [{"k": "v1"}, {"k": "v2"}]


def test_load_empty_sheet(tmp_path: Path):
    wb = Workbook()
    f = tmp_path / "v.xlsx"
    wb.save(f)
    assert load_rows(f) == []


def test_load_missing(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        load_rows(tmp_path / "sem.xlsx")
