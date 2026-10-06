# -*- coding: utf-8 -*-
"""Testes do _stage_data: copiar data_path -> record_dir/data.xlsx.

Contrato:
  - sem data_path      -> sem copy, retorna None
  - data_path ausente  -> retorna 3 (erro), sem copy
  - data_path == dst   -> nao copy (mesmo caminho), retorna None
  - data_path externo  -> copy p/ record_dir/data.xlsx, retorna None
  - data.xlsx ja existe-> sobrescreve
"""
from pathlib import Path

from openpyxl import Workbook

import macro_recorder.replay as replay


def _mk_xlsx(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["cpf"])
    ws.append(["123"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_stage_absent_is_noop(tmp_path: Path):
    rec = tmp_path / "rec"; rec.mkdir()
    assert replay._stage_data(rec, None, verbose=False) is None
    assert not (rec / "data.xlsx").exists()


def test_stage_missing_source_returns_3(tmp_path: Path):
    rec = tmp_path / "rec"; rec.mkdir()
    assert replay._stage_data(rec, str(tmp_path / "sem.xlsx"), verbose=False) == 3
    assert not (rec / "data.xlsx").exists()


def test_stage_same_path_is_noop(tmp_path: Path):
    rec = tmp_path / "rec"; rec.mkdir()
    src = _mk_xlsx(rec / "data.xlsx")  # já no lugar
    assert replay._stage_data(rec, str(rec / "data.xlsx"), verbose=False) is None


def test_stage_external_copies_in(tmp_path: Path):
    rec = tmp_path / "rec"; rec.mkdir()
    src = _mk_xlsx(tmp_path / "externo.xlsx")
    assert replay._stage_data(rec, str(src), verbose=False) is None
    dst = rec / "data.xlsx"
    assert dst.exists()
    from openpyxl import load_workbook
    ws = load_workbook(dst).active
    assert [c.value for c in ws[1]] == ["cpf"]


def test_stage_overwrites_existing(tmp_path: Path):
    rec = tmp_path / "rec"; rec.mkdir()
    (rec / "data.xlsx").write_text("antigo", encoding="utf-8")
    src = _mk_xlsx(tmp_path / "novo.xlsx")
    assert replay._stage_data(rec, str(src), verbose=False) is None
    assert (rec / "data.xlsx").stat().st_size != 6  # mais que "antigo"
