# -*- coding: utf-8 -*-
"""Excel de apoio (.xlsx) para replay parametrizado.

1 linha do .xlsx = 1 execução da macro; cada coluna = 1 placeholder
({nome}, {cpf}, ...). O replay lê data.xlsx e roda run(data_row=...) por linha.
Sem pandas — apenas openpyxl (no venv). pyperclip faz o colar (acentos OK +
o dado não fica como string no .py gerado nem arg do subprocess).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import List, Dict, Optional

from openpyxl import Workbook
from openpyxl.styles import Font

from macro_recorder.events import Step, Action

# {campo} — letras, números, _, -
_PLACEHOLDER_RE = re.compile(r"\{([A-Za-z0-9_-]+)\}")


def extract_placeholders(steps: List[Step]) -> List[str]:
    """Placeholders únicos (ordem de aparição) nos steps TYPE não-redacted."""
    seen: list[str] = []
    for s in steps:
        if s.action == Action.TYPE.value and s.text and not s.redacted:
            for m in _PLACEHOLDER_RE.finditer(s.text):
                if m.group(1) not in seen:
                    seen.append(m.group(1))
    return seen


def make_template(steps: List[Step], out_path: Path | str) -> Path:
    """Gera data.xlsx vazio com 1 coluna por placeholder.

    Sem placeholder → 1 coluna genérica ``campo1`` (user adiciona).
    """
    placeholders = extract_placeholders(steps)
    header = placeholders if placeholders else ["campo1"]
    wb = Workbook()
    ws = wb.active
    ws.title = "dados"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        ws.column_dimensions[cell.column_letter].width = 18
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def load_rows(xlsx_path: Path | str) -> List[Dict[str, str]]:
    """Lê .xlsx → list de dicts (header → valor str). Linhas vazias puladas."""
    path = Path(xlsx_path)
    if not path.exists():
        raise FileNotFoundError(f"Excel de apoio não encontrado: {path}")
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else f"col{i}"
                  for i, h in enumerate(next(rows))]
    except StopIteration:
        wb.close()
        return []
    out: list[dict[str, str]] = []
    for row in rows:
        vals = [row[i] if i < len(row) else None for i in range(len(header))]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        out.append({header[i]: (str(vals[i]) if vals[i] is not None else "")
                    for i in range(len(header))})
    wb.close()
    return out
